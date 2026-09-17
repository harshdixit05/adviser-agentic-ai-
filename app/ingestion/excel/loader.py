"""
Reads the source workbook into validated course records.

Parsing is deliberately separate from writing to the database: this module
returns records and a report, and decides nothing about persistence. That makes
the awkward part — nine inconsistently authored sheets — testable on its own.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import yaml

from app.db.models import ContentStatus, CourseLevel
from app.ingestion.excel import normalizers as norm


@dataclass
class CourseRecord:
    title: str
    domain: str
    level: CourseLevel
    duration_hours: float | None
    status: ContentStatus | None
    track: str | None
    faculty: str | None
    modules: list[dict]
    personas: list[str]
    source_sheet: str
    source_row: int

    @property
    def slug(self) -> str:
        return norm.slugify(f"{self.domain}-{self.title}")


@dataclass
class Report:
    """What parsed, what did not, and why — printed after every ingest."""

    courses: list[CourseRecord] = field(default_factory=list)
    skipped_sheets: list[str] = field(default_factory=list)
    banner_rows: int = 0
    empty_rows: int = 0
    rejected: list[tuple[str, int, str]] = field(default_factory=list)
    unmapped_status: set[str] = field(default_factory=set)
    missing_hours: list[str] = field(default_factory=list)
    missing_modules: list[str] = field(default_factory=list)
    missing_personas: list[str] = field(default_factory=list)

    def render(self) -> str:
        by_domain: dict[str, int] = {}
        by_level: dict[str, int] = {}
        for course in self.courses:
            by_domain[course.domain] = by_domain.get(course.domain, 0) + 1
            by_level[course.level.value] = by_level.get(course.level.value, 0) + 1

        lines = [
            "",
            "═" * 68,
            f"  INGESTION REPORT — {len(self.courses)} courses parsed",
            "═" * 68,
            "",
            "  By domain:",
        ]
        for domain, count in sorted(by_domain.items()):
            lines.append(f"    {domain:<30} {count:>3}")
        lines += ["", "  By level:"]
        for level in CourseLevel:
            if level.value in by_level:
                lines.append(f"    {level.value:<30} {by_level[level.value]:>3}")

        lines += [
            "",
            "  Rows not treated as courses:",
            f"    level banner rows                {self.banner_rows:>3}",
            f"    blank rows                       {self.empty_rows:>3}",
            f"    rejected (see below)             {len(self.rejected):>3}",
        ]
        if self.skipped_sheets:
            lines.append(f"    sheets skipped by config         {', '.join(self.skipped_sheets)}")

        if self.rejected:
            lines += ["", "  Rejected rows:"]
            for sheet, row, reason in self.rejected[:20]:
                lines.append(f"    {sheet}!row{row}: {reason}")
            if len(self.rejected) > 20:
                lines.append(f"    ... and {len(self.rejected) - 20} more")

        lines += ["", "  Gaps to fill (parsed, but the sheet had no value):"]
        lines.append(f"    no duration                      {len(self.missing_hours):>3}")
        lines.append(f"    no modules                       {len(self.missing_modules):>3}")
        lines.append(f"    no personas                      {len(self.missing_personas):>3}")
        if self.unmapped_status:
            values = ", ".join(sorted(self.unmapped_status))
            lines.append(f"    unrecognised status values       {values}")
            lines.append("      (left NULL rather than guessed — confirm what these mean)")

        lines += ["", "═" * 68, ""]
        return "\n".join(lines)


class WorkbookLoader:
    def __init__(self, workbook_path: Path, mapping_path: Path):
        self.workbook_path = workbook_path
        self.mapping = yaml.safe_load(mapping_path.read_text())
        self.header_marker = self.mapping["header_marker"]
        self.banner = re.compile(self.mapping["banner_pattern"], re.I)
        self.aliases: dict[str, list[str]] = self.mapping["aliases"]
        self.sheets: dict[str, dict] = self.mapping.get("sheets", {})

    # -- column resolution -------------------------------------------------

    def _resolve_columns(self, header: list) -> dict[str, int]:
        """Map canonical field -> column index by matching header aliases."""
        found: dict[str, int] = {}
        for index, cell in enumerate(header):
            key = norm.header_key(cell)
            if not key:
                continue
            for field_name, alias_list in self.aliases.items():
                if field_name in found:
                    continue
                if any(key == norm.header_key(alias) for alias in alias_list):
                    found[field_name] = index
                    break
        return found

    def _find_header(self, rows: list[list]) -> int | None:
        for index, row in enumerate(rows):
            if any(self.header_marker in norm.header_key(cell) for cell in row):
                return index
        return None

    # -- sheet parsing -----------------------------------------------------

    def load(self) -> Report:
        workbook = openpyxl.load_workbook(self.workbook_path, data_only=True)
        report = Report()

        for sheet_name in workbook.sheetnames:
            config = self.sheets.get(sheet_name, {})
            if config.get("skip"):
                report.skipped_sheets.append(sheet_name)
                continue
            self._load_sheet(workbook[sheet_name], sheet_name, config, report)

        return report

    def _load_sheet(self, sheet, sheet_name: str, config: dict, report: Report) -> None:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        header_index = self._find_header(rows)

        if header_index is None:
            report.rejected.append((sheet_name, 0, "no header row containing 'Course Title'"))
            return

        columns = self._resolve_columns(rows[header_index])
        if "title" not in columns or "level" not in columns:
            report.rejected.append(
                (sheet_name, header_index + 1, f"missing required columns, found {list(columns)}")
            )
            return

        default_domain = config.get("domain", sheet_name)

        for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            self._load_row(row, offset, sheet_name, columns, default_domain, report)

    def _load_row(
        self,
        row: list,
        row_number: int,
        sheet_name: str,
        columns: dict[str, int],
        default_domain: str,
        report: Report,
    ) -> None:
        def cell(field_name: str):
            index = columns.get(field_name)
            return row[index] if index is not None and index < len(row) else None

        title = norm.clean(cell("title"))

        if not title:
            if not any(norm.clean(value) for value in row):
                report.empty_rows += 1
            return

        # "DISCOVERY LEVEL" and "FLUENCY — 2 courses" are layout, not courses.
        populated = sum(1 for value in row if norm.clean(value))
        if self.banner.match(title) and populated <= 2:
            report.banner_rows += 1
            return

        level = norm.parse_level(cell("level"))
        if level is None:
            report.rejected.append(
                (sheet_name, row_number, f"unreadable level {norm.clean(cell('level'))!r} for {title[:40]!r}")
            )
            return

        raw_status = norm.clean(cell("status"))
        status = norm.parse_status(raw_status)
        if raw_status and status is None:
            report.unmapped_status.add(raw_status)

        record = CourseRecord(
            title=title,
            # The mapping is canonical: the Category column spells the same
            # domain differently across sheets, and some sheets omit it.
            domain=default_domain,
            level=level,
            duration_hours=norm.parse_hours(cell("hours")),
            status=status,
            track=norm.clean(cell("track")) or None,
            faculty=norm.clean(cell("faculty")) or None,
            modules=norm.parse_modules(cell("curriculum")),
            personas=norm.parse_personas(cell("personas")),
            source_sheet=sheet_name,
            source_row=row_number,
        )

        if record.duration_hours is None:
            report.missing_hours.append(record.title)
        if not record.modules:
            report.missing_modules.append(record.title)
        if not record.personas:
            report.missing_personas.append(record.title)

        report.courses.append(record)
